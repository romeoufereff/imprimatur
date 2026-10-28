#!/usr/bin/env python3
"""Run the eval suite's scripted graders.

`evals/evals.json` has defined nine evals and specified how to grade them —
"qa.py for designer outputs, seeded-violations.json recall/precision for the audit
eval, field-presence checks for narrative briefs" — since it was written. Nothing
computed any of it. The suite was a specification a human executed by hand, which
in practice meant it was not executed.

This runs the parts that can be graded deterministically, and sets up the parts
that need a model rather than pretending it can score them.

    run_evals.py audit                        # fully scripted; the fixture-rot check
    run_evals.py render                       # render-check regressions: seeded FAIL, fixed PASS
    run_evals.py cobrand                      # co-brand lockup: applied on copy, seeded FAILs
    run_evals.py bookkeeping                  # log_slide.py --deviation: one line per slide
    run_evals.py export                       # PPTX: inline runs, sup/sub, text over bg layers
    run_evals.py contrast                     # contrast over sibling gradient / solid layers
    run_evals.py diagram                      # check_diagram.py: seeded defects FAIL, fixed slide passes
    run_evals.py designer   <dir-of-slides>   # qa.py + per-eval assertions
    run_evals.py narrative  <briefs.md>       # SLIDE BRIEF field presence
    run_evals.py setup <id> <dir>             # stage an eval's working folder

Exit 0 when every scripted assertion passes.
"""

import argparse
import glob
import json
import os
import re
import shutil
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
PLUGIN = os.path.dirname(HERE)
EVALS = os.path.join(PLUGIN, "skills", "imprimatur", "evals")
sys.path.insert(0, HERE)
from ds_config import load  # noqa: E402

GREEN, RED, DIM, OFF = "\033[32m", "\033[31m", "\033[2m", "\033[0m"
if not sys.stdout.isatty():
    GREEN = RED = DIM = OFF = ""


def _run(script, *args):
    r = subprocess.run([sys.executable, os.path.join(HERE, script), *args],
                       capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr


def _fails(output):
    return [ln.strip()[6:].strip() for ln in output.splitlines()
            if ln.strip().startswith("FAIL")]


# ── audit: the one the eval contract names and nothing ever ran ──────────
def grade_audit():
    """Recall and precision against the pack's seeded-violation fixtures.

    Two things are being checked, and the second is the one that rots. Recall: does
    the script suite still flag every seed it is supposed to? Precision: does it flag
    anything *else*? A fixture that grows extra failures stops being ground truth —
    that already happened once, when stale font paths and a dropped data-template
    added six phantom failures and the eval had to be repaired rather than trusted.
    """
    ds = load()
    truth_path = os.path.join(ds.root, "evals", "stage-brand-audit", "seeded-violations.json")
    if not os.path.isfile(truth_path):
        print(f"{DIM}pack ships no seeded-violation fixtures — skipping{OFF}")
        return 0

    truth = json.load(open(truth_path, encoding="utf-8"))
    expected = truth.get("expected_script_output", {})
    ok = True

    print(f"audit fixtures  [{ds.name}]\n")
    total_seeds = script_detectable = flagged = false_pos = 0

    for fname, spec in truth["files"].items():
        path = os.path.join(ds.root, "evals", "stage-brand-audit", fname)
        if not os.path.isfile(path):
            print(f"  {RED}✗{OFF} {fname}: fixture missing")
            ok = False
            continue

        _, vout = _run("validate.py", path)
        _, cout = _run("check_contrast.py", path)
        v_fails, c_fails = _fails(vout), _fails(cout)

        want = expected.get(fname, {})
        v_want, c_want = want.get("validate.py"), want.get("check_contrast.py")

        seeds = spec["violations"]
        total_seeds += len(seeds)
        # Read the explicit flag, not the prose: one seed's `detectable_by` names
        # validate.py in order to say it does not catch it.
        scripted = [s for s in seeds if s.get("scripted") is True] or \
                   [s for s in seeds if s.get("scripted") is None
                    and ("validate.py" in s["detectable_by"]
                         or "check_contrast.py" in s["detectable_by"])]
        script_detectable += len(scripted)
        flagged += len(v_fails) + len(c_fails)

        print(f"  {fname}")
        print(f"    seeds: {len(seeds)}  ({len(scripted)} script-detectable, "
              f"{len(seeds) - len(scripted)} judgment-only)")

        for label, got, want_n in (("validate.py", len(v_fails), v_want),
                                   ("check_contrast.py", len(c_fails), c_want)):
            if want_n is None:
                print(f"    {DIM}{label}: {got} (no expectation recorded){OFF}")
                continue
            mark = f"{GREEN}✓{OFF}" if got == want_n else f"{RED}✗{OFF}"
            note = "" if got == want_n else \
                f"  {RED}— fixture rot: any failure not in the ground truth is a false positive{OFF}"
            print(f"    {mark} {label}: {got} failure(s), expected {want_n}{note}")
            if got != want_n:
                ok = False
                false_pos += max(0, got - want_n)
                for f in (v_fails if label == "validate.py" else c_fails):
                    print(f"        {f[:110]}")
        print()

    judgment = total_seeds - script_detectable
    print(f"  recall (scripted)  : {script_detectable}/{script_detectable} of the "
          f"script-detectable seeds" if ok else
          f"  recall (scripted)  : MISMATCH — see above")
    print(f"  precision          : {'no false positives' if not false_pos else f'{false_pos} false positive(s)'}")
    print(f"  {DIM}{judgment} seed(s) are judgment-only and are the brand-audit agent's job, "
          f"not a script's{OFF}")
    print(f"\n{GREEN if ok else RED}audit eval: {'PASS' if ok else 'FAIL'}{OFF}")
    return 0 if ok else 1


def _strings(node):
    """Every string value in a JSON node, joined — needles match raw text, not escaped JSON."""
    if isinstance(node, dict):
        return "\n".join(_strings(v) for v in node.values())
    if isinstance(node, list):
        return "\n".join(_strings(v) for v in node)
    return node if isinstance(node, str) else ""


# ── render: defects qa.py once passed ────────────────────────────────────
def grade_render():
    """Each case is a pair: a slide seeded with a layout defect a human reviewer caught
    after qa.py passed it, and the same slide fixed. The seeded one must FAIL with the
    expected finding kind, and its qa.py --json must name the offending element (the
    case's `needle`); the fixed one must pass with zero failures. A check that cannot
    fail its own repro is not a gate."""
    ds = load()
    root = os.path.join(ds.root, "evals", "render-checks")
    truth_path = os.path.join(root, "expected.json")
    if not os.path.isfile(truth_path):
        print(f"{DIM}pack ships no render-check fixtures — skipping{OFF}")
        return 0
    truth = json.load(open(truth_path, encoding="utf-8"))
    print(f"render-check fixtures  [{ds.name}]\n")
    ok = True

    def qa_json(path):
        r = subprocess.run([sys.executable, os.path.join(HERE, "qa.py"), "--files", path,
                            "--json", "--no-cache"], capture_output=True, text=True)
        try:
            return json.loads(r.stdout)
        except ValueError:
            return {"_error": (r.stdout + r.stderr)[-400:]}

    for case in truth["cases"]:
        print(f"  {case['id']}")
        bad = qa_json(os.path.join(root, case["bad"]))
        fixed = qa_json(os.path.join(root, case["fixed"]))
        checks = []
        if "_error" in bad or "_error" in fixed:
            checks.append(("qa.py --json ran", False))
            print(f"        {(bad.get('_error') or fixed.get('_error'))}")
        else:
            entry = bad["render_checks"].get(case["bad"], {})
            found = entry.get(case["section"]) or []
            kinds = {f.get("kind") for f in found}
            hits = [f for f in found if f.get("kind") == case["kind"]]
            checks += [
                (f"{case['bad']} fails", not bad["pass"]),
                (f"…with '{case['kind']}' in {case['section']}", bool(hits)),
                (f"…naming the element ({case['needle']!r})",
                 any(case["needle"] in _strings(f) for f in hits)),
            ]
            for extra in case.get("also") or []:
                checks.append((f"…and '{extra}'", extra in kinds))
            checks.append((f"{case['fixed']} passes (0 failures)",
                           fixed["pass"] and fixed["total_failures"] == 0))
        for label, passed in checks:
            print(f"    {GREEN + '✓' + OFF if passed else RED + '✗' + OFF} {label}")
            ok &= passed
        if not all(p for _, p in checks) and "_error" not in fixed and not fixed.get("pass"):
            for name, e in fixed.get("render_checks", {}).items():
                for sec in ("overflow", "collisions", "svg_sizing", "logo", "paint"):
                    for f in (e.get(sec) or [])[:3]:
                        print(f"        fixed/{sec}: {json.dumps(f, ensure_ascii=False)[:110]}")
            for name, v in fixed.get("validate", {}).items():
                for f in v["fails"][:3]:
                    print(f"        fixed/validate: {f[:110]}")
        print()
    print(f"{GREEN if ok else RED}render eval: {'PASS' if ok else 'FAIL'}{OFF}")
    return 0 if ok else 1


# ── cobrand: the lockup is applied at copy time and cannot silently go ─────
def _cobrand_templates(ds):
    """Pick templates from the ACTIVE pack by role: {'cover': [...], 'content': [...],
    'white': [...], 'closing': [...]} — content = dark-logo templates, white = white-logo
    ones, in name order, never a cover or closing template; closing = the pack's
    templateRoles.closing templates that carry a logo."""
    import cobrand
    roles = ds.get("templateRoles") or {}
    covers = list(roles.get("cover") or [])
    closings = list(roles.get("closing") or [])
    skip = set(covers) | set(closings)
    sel = (ds.get("logo") or {}).get("selector")
    out = {"cover": covers[:1], "content": [], "white": [], "closing": []}
    for p in sorted(glob.glob(os.path.join(ds.templates_dir, "*.html"))):
        stem = os.path.splitext(os.path.basename(p))[0]
        html = open(p, encoding="utf-8").read()
        span = cobrand.find_logo(html, sel)
        if stem in closings and span:
            out["closing"].append(stem)
        if stem in skip:
            continue
        if span:
            out["white" if cobrand._is_white(html[span[0]:span[0] + 600]) else "content"].append(stem)
    return out


def grade_cobrand():
    """The 1.6.1 co-brand contract, end to end on a temp deck staged from
    evals/stage-cobrand/: new_slide.py applies the declared lockup to every copy, a
    clean deck passes qa.py with no cobrand or logo findings, and each seeded mutation
    (removed, swapped, broken image, lifted, shrunk, white mark unfiltered) FAILs on its
    own slide with exactly one cobrand finding while the untouched slide stays clean."""
    import tempfile
    ds = load()
    src = os.path.join(EVALS, "stage-cobrand")
    truth = json.load(open(os.path.join(src, "seeded-mutations.json"), encoding="utf-8"))
    pool = _cobrand_templates(ds)
    print(f"co-brand fixture  [{ds.name}]\n")
    ok = True

    def mark(label, passed, detail=""):
        nonlocal ok
        ok &= bool(passed)
        print(f"    {GREEN + '✓' + OFF if passed else RED + '✗' + OFF} {label}"
              + (f"  {DIM}{detail[:160]}{OFF}" if detail and not passed else ""))

    def qa(deck):
        r = subprocess.run([sys.executable, os.path.join(HERE, "qa.py"), "--deck-dir", deck,
                            "--json", "--no-cache"], capture_output=True, text=True)
        try:
            return json.loads(r.stdout)
        except ValueError:
            return {"_error": (r.stdout + r.stderr)[-400:]}

    with tempfile.TemporaryDirectory() as tmp:
        deck = os.path.join(tmp, "deck")
        shutil.copytree(src, deck, ignore=shutil.ignore_patterns("seeded-mutations.json"))
        used, plan = {k: 0 for k in pool}, []
        for i, spec in enumerate(truth["slides"], 1):
            stems = pool.get(spec["role"]) or []
            if used[spec["role"]] >= len(stems):
                mark(f"pack has a {spec['role']} template for slide {i}", False)
                continue
            stem = stems[used[spec["role"]]]
            used[spec["role"]] += 1
            r = subprocess.run([sys.executable, os.path.join(HERE, "new_slide.py"), "--deck-dir", deck,
                                "--template", stem, "--n", str(i), "--slug", spec["role"]],
                               capture_output=True, text=True)
            name = f"{i:02d}-{spec['role']}.html"
            mark(f"{name} ({stem}): new_slide.py applies the lockup",
                 r.returncode == 0 and "cobrand: lockup applied" in r.stdout, r.stdout + r.stderr)
            plan.append((name, spec))
        print()

        clean = qa(deck)
        if "_error" in clean:
            mark("qa.py --json ran on the clean deck", False, clean["_error"])
            return 1
        rc = clean["render_checks"]
        noisy = {n: (e.get("cobrand") or []) + (e.get("cobrand_deck") or []) + (e.get("logo") or [])
                 for n, e in rc.items()}
        mark("clean deck: no cobrand or logo findings on any slide",
             not any(noisy.values()), json.dumps({n: f for n, f in noisy.items() if f})[:300])
        mark("clean deck: every slide measured (no lockup exempted or skipped)",
             all((e.get("cobrand_info") or {}).get("bottom") is not None for e in rc.values()))
        # 1.6.5: the client mark is sized from the pack logo beside it, never taller —
        # a closing template carrying the content-size mark got cover sizes (48 px client
        # wordmark next to a 28 px EPAM mark) and both scripts agreed it was fine
        tall = {n: (e.get("cobrand_info") or {}).get("measured") for n, e in rc.items()
                if ((e.get("cobrand_info") or {}).get("measured") or {}).get("client_h", 0)
                > ((e.get("cobrand_info") or {}).get("measured") or {}).get("logo_h", 0) + 0.5}
        mark("clean deck: no client mark taller than the pack logo it pairs with", not tall,
             json.dumps(tall))
        print()

        for name, spec in plan:
            for pat, repl in truth["mutations"].get(spec["mutation"] or "", []):
                path = os.path.join(deck, name)
                html = open(path, encoding="utf-8").read()
                new = re.sub(pat, repl, html, count=1)
                mark(f"{name}: seeded '{spec['mutation']}'", new != html)
                open(path, "w", encoding="utf-8").write(new)
        bad = qa(deck)
        if "_error" in bad:
            mark("qa.py --json ran on the seeded deck", False, bad["_error"])
            return 1
        mark("seeded deck: qa.py FAILs", not bad["pass"])
        for name, spec in plan:
            e = bad["render_checks"].get(name) or {}
            got = (e.get("cobrand") or []) + (e.get("cobrand_deck") or [])
            if spec["mutation"]:
                mark(f"{name}: exactly one cobrand FAIL with '{spec['expect']}'",
                     len(got) == 1 and spec["expect"] in got[0], json.dumps(got))
            else:
                mark(f"{name}: untouched slide stays clean", not got, json.dumps(got))
    print(f"\n{GREEN if ok else RED}cobrand eval: {'PASS' if ok else 'FAIL'}{OFF}")
    return 0 if ok else 1


# ── bookkeeping: a slide keeps one Deviations line ─────────────────────────
def grade_bookkeeping():
    """log_slide.py --deviation on a copy of evals/stage-bookkeeping/: replay the
    revision-round calls (replay.json) and require one Deviations line per slide, each
    note on its own slide's line, the earlier story kept under --append and replaced
    without it, and every line the calls did not target left byte-identical."""
    import tempfile
    src = os.path.join(EVALS, "stage-bookkeeping")
    calls = json.load(open(os.path.join(src, "replay.json"), encoding="utf-8"))["calls"]
    before = open(os.path.join(src, "design-decisions.md"), encoding="utf-8").read()
    print("bookkeeping fixture  [log_slide.py --deviation]\n")
    ok = True

    def mark(label, passed, detail=""):
        nonlocal ok
        ok &= bool(passed)
        print(f"    {GREEN + '✓' + OFF if passed else RED + '✗' + OFF} {label}"
              + (f"  {DIM}{detail[:200]}{OFF}" if detail and not passed else ""))

    def dev_lines(text, n):
        return [l for l in text.splitlines() if re.match(rf"^- 0*{n}:", l)]

    with tempfile.TemporaryDirectory() as deck:
        shutil.copy2(os.path.join(src, "design-decisions.md"), deck)
        for c in calls:
            args = ["--deck-dir", deck, "--n", str(c["n"]), "--file", c["file"], "--template",
                    c["template"], "--visual", c["visual"], "--focal", c["focal"], "--status",
                    c["status"], "--deviation", c["deviation"]] + (["--append"] if c["append"] else [])
            rc, out = _run("log_slide.py", *args)
            mark(f"slide {c['n']:02d}: log_slide.py --deviation{' --append' if c['append'] else ''} exits 0",
                 rc == 0, out)
        after = open(os.path.join(deck, "design-decisions.md"), encoding="utf-8").read()
    print()
    targeted = {c["n"] for c in calls}
    for c in calls:
        n, old = c["n"], dev_lines(before, c["n"])
        new = dev_lines(after, n)
        mark(f"slide {n:02d}: exactly one Deviations line", len(new) == 1, str(new))
        if len(new) != 1:
            continue
        mark(f"slide {n:02d}: the replayed note is on that line", c["deviation"] in new[0], new[0])
        if old and c["append"]:
            story = old[-1].split(":", 1)[1].strip()
            mark(f"slide {n:02d}: --append kept the slide's earlier story on the same line",
                 new[0].startswith(f"- {n:02d}: {story}"), new[0])
        elif old:
            mark(f"slide {n:02d}: re-log without --append replaced the earlier story",
                 old[-1] not in after, new[0])
    untouched = lambda text: [l for l in text.splitlines()
                              if not re.match(r"^(- |\| )0*(%s)[:|\s]" % "|".join(map(str, targeted)), l)]
    mark("every line the calls did not target is byte-identical", untouched(before) == untouched(after))
    print(f"\n{GREEN if ok else RED}bookkeeping eval: {'PASS' if ok else 'FAIL'}{OFF}")
    return 0 if ok else 1


# ── export: what the PPTX exporter makes of inline runs and background layers ──
EXPORT_FIXTURES = os.path.join(PLUGIN, "skills", "pptx-export", "tests", "fixtures")


def _export_fixtures(tmp):
    """extract_ir + build_pptx on the export fixtures → {name: (Presentation, IR)}."""
    sys.path.insert(0, os.path.join(PLUGIN, "skills", "pptx-export", "scripts"))
    from playwright.sync_api import sync_playwright
    from pptx import Presentation
    from extract_ir import extract
    from build_pptx import build
    out = {}
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": 1920, "height": 1080})
        for name in ("inline-runs", "bg-layers", "img-svg"):
            ir = extract(page, os.path.join(EXPORT_FIXTURES, f"{name}.html"), tmp, name)
            ir_path = os.path.join(tmp, f"{name}.json")
            json.dump(ir, open(ir_path, "w", encoding="utf-8"))
            build([ir_path], os.path.join(tmp, f"{name}.pptx"))
            out[name] = (Presentation(os.path.join(tmp, f"{name}.pptx")), ir)
        browser.close()
    return out


def grade_export():
    """skills/pptx-export/tests/fixtures/, exported for real and read back with python-pptx.
    inline-runs: a styled run inside a word gets no space (`mig<span>V</span>isor` exported
    as "mig V isor"), <sup>/<sub> are baseline-offset runs on the same line (™ became a
    paragraph of its own), real spaces, soft wraps and <br> survive. bg-layers: text on a
    background-image element stays editable text over a picture of the background alone
    (the gantt track and the gradient-border card used to be one screenshot, text
    included), and a `to right` gradient without stop positions is a native gradient fill.
    img-svg (1.6.6): an <img> of an SVG file or data: URI is native shapes in its own class
    colours, white under brightness(0) invert(1), and stays a picture under any other filter
    (the co-brand partner logo shipped as a 1x PNG)."""
    import tempfile
    from pptx.oxml.ns import qn
    print("export fixtures  [pptx-export]\n")
    ok = True

    def mark(label, passed, detail=""):
        nonlocal ok
        ok &= bool(passed)
        print(f"    {GREEN + '✓' + OFF if passed else RED + '✗' + OFF} {label}"
              + (f"  {DIM}{str(detail)[:220]}{OFF}" if detail and not passed else ""))

    def frames(prs):
        return [sh for sh in prs.slides[0].shapes if sh.has_text_frame and sh.text_frame.text.strip()]

    def find(prs, needle):
        return next((sh for sh in frames(prs) if needle in sh.text_frame.text.replace("\n", " ")
                     .replace("\x0b", " ")), None)

    def paras(sh):
        return ["".join(r.text for r in p.runs) for p in sh.text_frame.paragraphs]

    def baselines(sh):
        return [(r.text, r.font._rPr.get("baseline")) for p in sh.text_frame.paragraphs for r in p.runs]

    with tempfile.TemporaryDirectory() as tmp:
        try:
            got = _export_fixtures(tmp)
        except Exception as e:  # noqa: BLE001 — report, don't crash the suite
            mark("export fixtures extract + build", False, repr(e))
            return 1
        prs, _ = got["inline-runs"]
        print("  inline-runs.html")
        sh = find(prs, "mig")
        mark("wordmark: one paragraph reading 'migVisor™'", sh is not None and paras(sh) == ["migVisor™"],
             sh and paras(sh))
        mark("wordmark: ™ is a superscript run (baseline +30 %), no run starts with a space",
             sh is not None and ("™", "30000") in baselines(sh)
             and not any(t.startswith(" ") for t, _ in baselines(sh)), sh and baselines(sh))
        sh = find(prs, "Brownfield")
        mark("real spaces around a styled run are kept", sh is not None
             and paras(sh) == ["Brownfield intake stays editable"], sh and paras(sh))
        sh = find(prs, "and E = mc")
        mark("H₂O / mc²: one paragraph, sub −25 % and sup +30 % runs", sh is not None
             and paras(sh) == ["H2O and E = mc2"] and ("2", "-25000") in baselines(sh)
             and ("2", "30000") in baselines(sh), sh and baselines(sh))
        sh = find(prs, "larger")
        mark("a larger inline span stays on its line", sh is not None
             and paras(sh) == ["One larger word on one line"], sh and paras(sh))
        sh = find(prs, "A long sentence")
        mark("a soft-wrapped sentence keeps its line breaks and single spaces", sh is not None
             and len(paras(sh)) >= 2 and " ".join(paras(sh)) == "A long sentence that the browser "
             "wraps onto more than one line inside its narrow box", sh and paras(sh))
        sh = find(prs, "Before the break")
        mark("<br> is still a line break", sh is not None
             and paras(sh) == ["Before the break", "after the break"], sh and paras(sh))

        prs, ir = got["bg-layers"]
        print("\n  bg-layers.html")
        for needle in ("MVP until end of the year", "Native gradient pill", "Card heading",
                       "Label over a sibling gradient"):
            mark(f"'{needle}' is editable text", find(prs, needle) is not None)
        bg_only = {n["cls"]: n for n in ir["nodes"] if n.get("bgOnly")}
        mark("track and card export a background-only picture", set(bg_only) == {"track", "card"},
             sorted(bg_only))
        pics = [n for n in ir["nodes"] if n["kind"] == "raster"]
        mark("the content-free radial wash is still one picture, and nothing else is",
             sorted(n["cls"] for n in pics) == ["card", "track", "wash"], [n["cls"] for n in pics])
        try:
            from PIL import Image
            for cls, inset in (("track", 0), ("card", 20)):
                node = bg_only.get(cls) or {}
                im = Image.open(node["asset"]).convert("RGBA")
                w, h = im.size
                px = im.load()
                inked = sum(1 for x in range(inset, w - inset, 2) for y in range(inset, h - inset, 2)
                            if px[x, y][3] > 128 and max(px[x, y][:3]) - min(px[x, y][:3]) > 60
                            or (px[x, y][3] > 128 and max(px[x, y][:3]) < 120))
                mark(f"{cls}: the picture holds no content pixels (bar / heading hidden for the shot)",
                     inked == 0, f"{inked} inked samples")
        except ImportError:
            print(f"    {DIM}Pillow missing — picture-content rows skipped{OFF}")
        pill = find(prs, "Native gradient pill")
        pill_fill = [sh for sh in prs.slides[0].shapes if pill is not None and sh is not pill
                     and sh.element.find(".//" + qn("a:gradFill")) is not None
                     and abs(sh.left - pill.left) < 20000 and abs(sh.top - pill.top) < 20000]
        mark("`to right` gradient without stop positions is a native gradient fill under the pill",
             bool(pill_fill))

        prs, ir = got["img-svg"]
        print("\n  img-svg.html")
        from build_pptx import px as to_emu
        imgs = {n["cls"]: n for n in ir["nodes"] if n["kind"] == "raster" and n["tag"] == "img"}
        mark("four <img> nodes, each flagged as an SVG source", sorted(imgs) == [
             "mark-data", "mark-plain", "mark-shadow", "mark-white"] and all(
             n.get("imgSvg") for n in imgs.values()), sorted(imgs))

        def at(node):
            r = node["rect"]
            return next((sh for sh in prs.slides[0].shapes if abs(sh.left - to_emu(r["x"])) < 13000
                         and abs(sh.top - to_emu(r["y"])) < 13000), None)

        def colours(sh):
            return {c.get("val").upper() for c in sh.element.iter(qn("a:srgbClr"))}

        rgb = lambda *c: "%02X%02X%02X" % c  # noqa: E731 — fixture colours, not brand values
        cases = (("mark-plain", 2, {rgb(51, 51, 51), rgb(48, 85, 218)}),
                 ("mark-white", 2, {rgb(255, 255, 255)}),
                 ("mark-data", 1, {rgb(48, 85, 218)}))
        for cls, n_shapes, want in cases:
            node = imgs.get(cls) or {"rect": {"x": -1e6, "y": -1e6}}
            sh = at(node)
            is_group = sh is not None and sh.shape_type == 6
            got_n = len(sh.shapes) if is_group else None
            mark(f"{cls}: a native group of {n_shapes} shape(s), no picture", is_group
                 and got_n == n_shapes, sh is not None and (sh.shape_type, got_n))
            mark(f"{cls}: paints {sorted(want)} (the file's <style>, the page's .st0 kept out)",
                 is_group and colours(sh) == want, is_group and sorted(colours(sh)))
        sh = at(imgs.get("mark-shadow") or {"rect": {"x": -1e6, "y": -1e6}})
        mark("mark-shadow: a drop-shadow filter keeps the picture",
             sh is not None and sh.shape_type == 13 and not (imgs.get("mark-shadow") or {}).get("asset_svg"),
             sh is not None and sh.shape_type)
    print(f"\n{GREEN if ok else RED}export eval: {'PASS' if ok else 'FAIL'}{OFF}")
    return 0 if ok else 1


def grade_contrast():
    """check_contrast + check_gradient on bg-layers.html: text over an absolutely positioned
    SIBLING gradient layer is deferred to the pixel sample (it read as white-on-white,
    1.00:1 FAIL, unless the parent got a solid colour), and text over a pale sibling SOLID
    layer FAILs with that layer's real ratio, not the white page's."""
    import check_contrast
    import check_gradient
    from playwright.sync_api import sync_playwright
    print("contrast fixture  [bg-layers.html]\n")
    ok = True

    def mark(label, passed, detail=""):
        nonlocal ok
        ok &= bool(passed)
        print(f"    {GREEN + '✓' + OFF if passed else RED + '✗' + OFF} {label}"
              + (f"  {DIM}{str(detail)[:220]}{OFF}" if detail and not passed else ""))

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": 1920, "height": 1080})
        page.goto("file://" + os.path.join(EXPORT_FIXTURES, "bg-layers.html"), wait_until="networkidle")
        fails, manual = check_contrast.analyze_page(page)
        grad = check_gradient.analyze_page(page)
        browser.close()
    by_text = {it["text"]: (ratio, need) for it, ratio, need in (fails or [])}
    mark("label over a sibling gradient layer is not a computed FAIL",
         not any(t.startswith("Label over a sibling") for t in by_text), by_text)
    sib = next((g for g in grad if g.get("text", "").startswith("Label over a sibling")), None)
    mark("…it is pixel-sampled by check_gradient, worst case above 1.5:1 (not white-on-white)",
         sib is not None and not sib.get("unsampled") and sib.get("ratio_worst", 0) > 1.5, sib)
    low = next((v for t, v in by_text.items() if t.startswith("White on a pale")), None)
    mark("white text on a pale sibling solid layer FAILs with that layer's ratio (1.05–1.3:1)",
         low is not None and 1.05 < low[0] < 1.3, low)
    mark("text inside the gradient pill / track / card is deferred, not failed",
         not any(t.startswith(("MVP", "Native gradient", "Card heading")) for t in by_text), by_text)
    print(f"\n{GREEN if ok else RED}contrast eval: {'PASS' if ok else 'FAIL'}{OFF}")
    return 0 if ok else 1


# ── diagram: check_diagram.py and the QA freeze (2026-10-26 restyle review) ─────
DIAGRAM_FIXTURES = os.path.join(EVALS, "stage-diagram")


def grade_diagram():
    """seeded.html carries one of each defect a 2026-10-26 restyle review found only after a
    script measured it, plus the look-alikes that must NOT fire (a top rule is not a box, a
    tight badge numeral, screen-reader-only text off the canvas, a line masked by a
    white-backed label, an exit arrow with nothing ahead). fixed.html repairs each defect and
    must report nothing. The seeded set includes the stripe card (a 3 px coloured top edge over
    a 1 px border) two critique rounds approved and the owner rejected, and the "one short
    card" a grid row leaves when it stretches the wrapper, not the card. One card animates with a deck-style `!important` rule: it must
    spill without render_checks' freeze and sit inside its card with it. vocabulary*.html
    grade the panel-style / tint / connector-colour rows against role colours passed in."""
    import check_diagram as cd
    import check_overflow
    import render_checks
    from playwright.sync_api import sync_playwright
    ds = load()
    print("diagram fixtures  [stage-diagram]\n")
    ok = True

    def mark(label, passed, detail=""):
        nonlocal ok
        ok &= bool(passed)
        print(f"    {GREEN + '✓' + OFF if passed else RED + '✗' + OFF} {label}"
              + (f"  {DIM}{str(detail)[:240]}{OFF}" if detail and not passed else ""))

    cfg = cd.config(ds)
    cfg["banBars"] = True   # the fixture grades the check, not the active pack's toggle
    voc_cfg = dict(cfg, tint=["#F5F5F5"], connector="#3355DD", exception="#6A6A6A",
                   decision=["#B4690E", "#E6E8EC"], maxPanelStyles=3, maxTintPanels=1)

    def run(page, name, freeze, c=cfg):
        page.goto("file://" + os.path.join(DIAGRAM_FIXTURES, name), wait_until="networkidle")
        page.wait_for_timeout(400)
        if freeze:
            page.evaluate(render_checks.FREEZE_JS, render_checks.FREEZE_CSS)
        return page.evaluate(cd.DIAGRAM_JS, [c])

    def has(found, kind, needle=""):
        return any(f["kind"] == kind and needle in json.dumps(f, ensure_ascii=False) for f in found)

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": 1920, "height": 1080})
        print("  seeded.html (frozen)")
        res = run(page, "seeded.html", True)
        found = res["findings"]
        mark("text leaves its box: the fixed-height card", has(found, "text leaves its box", "Card copy"), found)
        mark("connector through text: a route starting in a descender band",
             has(found, "connector through text", "Descender gyp label"), found)
        mark("arrow tip buried inside Target one", has(found, "arrow tip buried", "Target one"), found)
        mark("arrow tip floats short of Target two", has(found, "arrow tip floats", "Target two"), found)
        mark("connector contrast: rule-grey exception path", has(found, "connector contrast"), found)
        mark("accent stripe on the left (an accent bar)", has(found, "accent stripe", "Accent bar note"), found)
        mark("accent stripe on top (the stripe card the owner rejected)", has(found, "accent stripe", "Stripe card copy"), found)
        mark("uneven row: the card that did not stretch with its wrapper",
             has(found, "uneven row", "Short row card"), found)
        quiet = [f for f in found if any(n in json.dumps(f) for n in
                 ("Masked label", "Text under a top rule", '"07"', "Animated card copy", "1600", "Divider above"))]
        mark("no finding on the look-alikes (top rule, divider, badge, masked line, exit arrow, frozen card)",
             not quiet, quiet)
        mark("exactly the eight seeded findings", len(found) == 8, [cd.describe(f) for f in found])
        ov, _ = check_overflow.analyze_overflow_page(page, ds, "seeded.html")
        cols, _ = check_overflow.analyze_collisions_page(page, 4, "")
        mark("screen-reader-only text is not an overflow or a collision",
             not any("Screen reader" in json.dumps(x) for x in (ov or []) + (cols or [])), (ov, cols))
        print("  seeded.html (not frozen)")
        found = run(page, "seeded.html", False)["findings"]
        mark("the mid-animation card spills without the freeze (the freeze is load-bearing)",
             has(found, "text leaves its box", "Animated card copy"), found)
        print("  fixed.html (frozen)")
        found = run(page, "fixed.html", True)["findings"]
        mark("zero findings", not found, [cd.describe(f) for f in found])
        print("  vocabulary.html")
        voc = run(page, "vocabulary.html", True, voc_cfg)["vocabulary"]
        kinds = {f["kind"] for f in cd.vocabulary_findings(voc, voc_cfg)}
        want = {"too many panel styles", "too many tinted panels", "connector colours", "exception path colour"}
        mark("four vocabulary rows fire", kinds == want, (kinds, voc))
        mark("the decision tint is found on the decision card",
             any("A person decides" in d for d in voc["decision_colour"]), voc["decision_colour"])
        print("  vocabulary-fixed.html")
        voc = run(page, "vocabulary-fixed.html", True, voc_cfg)["vocabulary"]
        rows = cd.vocabulary_findings(voc, voc_cfg)
        mark("no vocabulary rows; three panel styles (which tint fills a card is meaning, not style)",
             not rows and voc["panel_styles"] == 3, (rows, voc["panels"]))
        browser.close()
    print(f"\n{GREEN if ok else RED}diagram eval: {'PASS' if ok else 'FAIL'}{OFF}")
    return 0 if ok else 1


# ── designer: grade produced slides ──────────────────────────────────────
def grade_designer(target):
    slides = sorted(glob.glob(os.path.join(target, "[0-9]*.html"))) \
        if os.path.isdir(target) else [target]
    if not slides:
        print(f"{RED}no NN-*.html slides in {target}{OFF}")
        return 1

    print(f"designer outputs ({len(slides)} slide(s))\n")
    ok = True
    for s in slides:
        name = os.path.basename(s)
        rc, out = _run("qa.py", s)
        html = open(s, encoding="utf-8").read()

        checks = [
            ("mechanical QA (validate + contrast + overflow + paint)", rc == 0),
            ("carries data-template", bool(re.search(r'data-template="[^"]+"', html))),
            ("self-contained (no external css/js beyond CDNs)",
             not re.search(r'<link[^>]+href="(?!https?://)[^"]+\.css"', html)),
            ("no raster referenced by path (data: URIs only)",
             not re.search(r'<img[^>]+src="(?!data:)[^"]+"', html)),
        ]
        print(f"  {name}")
        for label, passed in checks:
            print(f"    {GREEN + '✓' + OFF if passed else RED + '✗' + OFF} {label}")
            ok &= passed
        if rc != 0:
            for f in _fails(out)[:6]:
                print(f"        {f[:110]}")
        print()
    print(f"{GREEN if ok else RED}designer eval: {'PASS' if ok else 'FAIL'}{OFF}")
    return 0 if ok else 1


# ── narrative: SLIDE BRIEF field presence ────────────────────────────────
REQUIRED_FIELDS = ("Message:", "Structure:", "Visual:", "Key data:",
                   "Emphasis:", "Audience:", "Density:")


def grade_narrative(path):
    text = open(path, encoding="utf-8").read()
    briefs = re.split(r'^SLIDE \d+:', text, flags=re.M)[1:]
    print(f"narrative briefs ({len(briefs)} found)\n")
    if not briefs:
        print(f"{RED}no 'SLIDE N:' blocks — the designer cannot consume this{OFF}")
        return 1

    ok = True
    for i, b in enumerate(briefs, 1):
        missing = [f for f in REQUIRED_FIELDS if f not in b]
        mark = f"{GREEN}✓{OFF}" if not missing else f"{RED}✗{OFF}"
        print(f"  {mark} slide {i}" + (f"  missing: {', '.join(missing)}" if missing else ""))
        ok &= not missing

    # The variance dial's whole purpose: enough slides carrying a real visual.
    visuals = [b for b in briefs if re.search(r'Visual:\s*(?!none)\S', b)]
    print(f"\n  visual slides: {len(visuals)}/{len(briefs)}"
          f"  {DIM}(taste-dials.md: low >=1, medium >=2, high >=3 incl. >=1 bespoke){OFF}")
    if len(visuals) == 0:
        print(f"  {RED}✗ an all-cards deck is the wall-of-cards tell{OFF}")
        ok = False
    print(f"\n{GREEN if ok else RED}narrative eval: {'PASS' if ok else 'FAIL'}{OFF}")
    return 0 if ok else 1


# ── setup: stage an eval's isolated working folder ───────────────────────
def setup(eval_id, dest):
    data = json.load(open(os.path.join(EVALS, "evals.json"), encoding="utf-8"))
    ev = next((e for e in data["evals"] if str(e["id"]) == str(eval_id)), None)
    if not ev:
        print(f"no eval with id {eval_id}. Available: "
              f"{', '.join(str(e['id']) for e in data['evals'])}")
        return 1
    os.makedirs(dest, exist_ok=True)
    for rel in ev.get("files") or []:
        src = os.path.join(EVALS, rel)
        if not os.path.isfile(src):
            print(f"  {RED}missing fixture: {rel}{OFF}")
            return 1
        shutil.copy2(src, os.path.join(dest, os.path.basename(rel)))
        print(f"  copied {os.path.basename(rel)}")
    print(f"\neval {ev['id']} — {ev['name']}\n  {dest}\n")
    print("PROMPT:")
    print("  " + (ev.get("prompt", "") or "").replace("\n", "\n  "))
    print("\nEXPECTED:")
    print("  " + (ev.get("expected_output", "") or "").replace("\n", "\n  "))
    print(f"\nGrade with:  run_evals.py designer {dest}")
    return 0


def main():
    ap = argparse.ArgumentParser(description=__doc__.strip().split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("audit", help="Scripted recall/precision on the pack's seeded fixtures")
    sub.add_parser("render", help="Render-check regressions: seeded slide FAILs, fixed slide passes")
    sub.add_parser("cobrand", help="Co-brand lockup: applied by new_slide.py, seeded mutations FAIL")
    sub.add_parser("bookkeeping", help="log_slide.py --deviation: one Deviations line per slide")
    sub.add_parser("export", help="PPTX export: inline runs, sup/sub, text over background layers")
    sub.add_parser("contrast", help="Contrast over sibling gradient / solid layers")
    sub.add_parser("diagram", help="Diagram checks: seeded defects FAIL, look-alikes and fixed slide pass")
    d = sub.add_parser("designer", help="Grade produced slides"); d.add_argument("target")
    n = sub.add_parser("narrative", help="Grade produced SLIDE BRIEFs"); n.add_argument("path")
    s = sub.add_parser("setup", help="Stage an eval's working folder")
    s.add_argument("eval_id"); s.add_argument("dest")
    a = ap.parse_args()

    if a.cmd == "audit":
        return grade_audit()
    if a.cmd == "render":
        return grade_render()
    if a.cmd == "cobrand":
        return grade_cobrand()
    if a.cmd == "bookkeeping":
        return grade_bookkeeping()
    if a.cmd == "export":
        return grade_export()
    if a.cmd == "contrast":
        return grade_contrast()
    if a.cmd == "diagram":
        return grade_diagram()
    if a.cmd == "designer":
        return grade_designer(a.target)
    if a.cmd == "narrative":
        return grade_narrative(a.path)
    return setup(a.eval_id, a.dest)


if __name__ == "__main__":
    sys.exit(main())
