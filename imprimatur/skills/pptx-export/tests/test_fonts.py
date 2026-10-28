"""Font-family resolution in the PPTX builder — run from validate_all.sh (pytest, no browser).

  two-family pack      a run keeps its computed family when the pack declares it: a
                       client pack (Century Gothic display + Arial text) exported every
                       title as Arial because the IR carried no family and the builder
                       had one name per weight band (2026-10-22).
  single-family pack   per-weight full names (Museo Sans 300/500/700) are untouched and
                       never get a bold flag, so such packs export byte-identically.
  shared-name bands    {"400": "Arial", "700": "Arial"}: bold is the only way to carry
                       the weight, so 700 sets it (bold Arial exported regular before).

The pack settings are patched onto the module, so the result does not depend on which
pack is active on this machine.
"""
import os
import sys

import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "scripts"))
import build_pptx  # noqa: E402

CG = '"Century Gothic", Arial, sans-serif'
ARIAL = "Arial, Helvetica, sans-serif"


@pytest.fixture
def two_family(monkeypatch):
    monkeypatch.setattr(build_pptx, "_FONT_BY_WEIGHT", [(700, "Arial"), (400, "Arial")])
    monkeypatch.setattr(build_pptx, "_FAMILY_PATTERN", "Century Gothic|Arial")


@pytest.fixture
def per_weight(monkeypatch):
    monkeypatch.setattr(build_pptx, "_FONT_BY_WEIGHT",
                        [(700, "Museo Sans 700"), (400, "Museo Sans 500"), (300, "Museo Sans 300")])
    monkeypatch.setattr(build_pptx, "_FAMILY_PATTERN", "Museo Sans")


def test_two_family_display_run_keeps_its_face(two_family):
    assert build_pptx.font_for(400, CG) == ("Century Gothic", False)
    assert build_pptx.font_for(700, CG) == ("Century Gothic", True)


def test_two_family_text_run_and_shared_name_bold(two_family):
    assert build_pptx.font_for(400, ARIAL) == ("Arial", False)
    assert build_pptx.font_for(700, ARIAL) == ("Arial", True)


def test_undeclared_or_missing_stack_falls_back_per_weight(two_family):
    assert build_pptx.font_for(400, "Georgia, serif") == ("Arial", False)
    assert build_pptx.font_for(400, None) == ("Arial", False)


def test_per_weight_names_unchanged_and_never_bold(per_weight):
    for w in (300, 400, 700):
        assert build_pptx.font_for(w, '"Museo Sans", sans-serif') == (build_pptx.font_for_weight(w), False)
        assert build_pptx.font_for(w, None) == (build_pptx.font_for_weight(w), False)


def test_text_box_writes_each_runs_family(two_family):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    st = {"size": 40, "weight": 700, "color": "#000000", "lineHeight": 48}
    node = {"kind": "text", "rect": {"x": 0, "y": 0, "w": 800, "h": 120},
            "align": "left", "anchor": "top", "insets": {"l": 0, "t": 0, "r": 0, "b": 0},
            "runs": [{"text": "Title", "st": dict(st, family=CG)},
                     {"text": "\n", "st": dict(st, family=CG)},
                     {"text": "Body", "st": dict(st, weight=400, family=ARIAL)}]}
    build_pptx.add_text(slide, node)
    runs = [r for p in slide.shapes[-1].text_frame.paragraphs for r in p.runs]
    assert [(r.text, r.font.name, r.font.bold) for r in runs] == [
        ("Title", "Century Gothic", True), ("Body", "Arial", False)]
